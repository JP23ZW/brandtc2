from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path
from functools import lru_cache

from openpyxl import load_workbook


# Ordered, conservative links between the standard measure wording and the
# corresponding row in Eenheidsprijzenlijst 2026.xlsx. A missing match is safer
# than assigning an unrelated price; inspectors can always override suggestions.
PRICE_RULES: list[tuple[str, list[int]]] = [
    (r"draadglas", [1]),
    (r"naden rondom (het )?kozijn", [91]),
    (r"glas in (het kozijn|de deur)|zijlicht|bovenlicht|brandwerend glas", [3]),
    (r"glazen bouwstenen|aantonen of .*constructieonderdeel|eigenaar moet aantonen", [101]),
    (r"vrijloopdranger.*aansturing", [12, 13]),
    (r"deurdranger.*toepassen", [6]),
    (r"deurdranger.*herstel", [104]),
    (r"deurdranger.*afstel|sluitkracht.*dranger", [11]),
    (r"dagschoot", [82]),
    (r"slotplaat|slot herstellen", [81]),
    (r"naden\s*<\s*3|naden kleiner dan 3", [18]),
    (r"ondernaad\s*<\s*6|ondernaad kleiner dan 6", [19]),
    (r"sponning", [21]),
    (r"deurnaald", [22]),
    (r"schroeven.*250", [24]),
    (r"opschuimende strips", [20]),
    (r"deur zonder kozijn vervangen", [26]),
    (r"deur met kozijn vervangen|hardhouten kozijn", [29]),
    (r"stootplaat|schopplaat", [23]),
    (r"koof.*30", [35]),
    (r"koof", [34]),
    (r"naden rondom schacht|naden.*afdichten", [91]),
    (r"schacht.*30", [37]),
    (r"schacht", [36]),
    (r"doorvoer", [42]),
    (r"sparing", [96]),
    (r"ventilatieventiel", [50]),
    (r"ventilatierooster", [49]),
    (r"brandklep", [44]),
    (r"putty", [100]),
    (r"knopcilinder", [86]),
    (r"begaanbaarheid.*vluchtweg|tegels herstellen", [51]),
    (r"bewoners informeren.*objecten.*vluchtweg", [52]),
    (r"trap in de vluchtweg|vluchttrap", [53]),
    (r"hellingbaan|hoogteverschil", [54]),
    (r"brandveilige scootmobiel|scootmobielen.*veilig gestald", [59]),
    (r"bewoners informeren.*scootmobiel", [58]),
    (r"oplaadpunt", [60]),
    (r"draairichting", [55]),
    (r"zelfsluitende constructieonderdeel.*niet vast", [56]),
    (r"blusmiddelen.*onderhouden|blusmiddelen.*keuringssticker", [61]),
    (r"voldoende.*blusmiddelen", [62]),
    (r"blusmiddel.*zichtbaar|blusmiddelen.*pictogram", [63]),
    (r"bmi.*nominale staat|bmi.*certific", [64]),
    (r"documenten.*bmi", [66]),
    (r"verlichting.*onderhouden|verlichting herstellen", [69]),
    (r"nood.*vluchtwegverlichting.*aanbrengen", [70]),
    (r"vluchtrouteaanduid", [71]),
    (r"inventariseren.*rookmelders", [76, 77]),
    (r"rookmelders plaatsen|rookmelder plaatsen", [77]),
    (r"bewoners informeren.*rookmelders", [80]),
    (r"brandweeringang", [74]),
    (r"dakconstructie.*control", [88]),
    (r"wand in de brandscheiding|brandscheidingen brandwerend", [107]),
]


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = value.lower().replace("besluit bouwwerken leefomgeving", "bbl")
    value = re.sub(r"\bminuten\b", "min", value)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9<>]+", " ", value)).strip()


def load_unit_prices(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["huidige versie"]
        result = []
        for row_number, (description, price, unit) in enumerate(
            sheet.iter_rows(min_row=1, max_col=3, values_only=True), start=1
        ):
            if not description:
                continue
            result.append(
                {
                    "id": row_number,
                    "description": str(description).strip(),
                    "price": float(price or 0),
                    "unit": str(unit or "").strip(),
                }
            )
        return result
    finally:
        workbook.close()


def parse_price_ids(value) -> list[int]:
    if isinstance(value, (list, dict)):
        raw = value
    else:
        try:
            raw = json.loads(value or "[]")
        except (TypeError, json.JSONDecodeError):
            raw = []
    if isinstance(raw, dict):
        raw = raw.get('items', [])
    if not isinstance(raw, list):
        return []
    result = []
    for item in raw:
        if isinstance(item, dict):
            item = item.get('id')
        try:
            number = int(item)
        except (TypeError, ValueError):
            continue
        if number not in result:
            result.append(number)
    return result


@lru_cache(maxsize=1)
def _rule_sources():
    path = Path(__file__).parent / 'data' / 'price_rule_sources.json'
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}


def _rule_prices(prices):
    """Resolve rule references by description/unit even if source rows move."""
    indexed = {}
    for price in prices:
        key = (normalize_text(price['description']), price.get('unit', ''))
        indexed.setdefault(key, []).append(price)
    result = {}
    for number, source in _rule_sources().items():
        matches = indexed.get((normalize_text(source['description']), source['unit']), [])
        if len(matches) == 1:
            result[int(number)] = matches[0]
    return result


def suggest_price_ids(measure: str, prices: list[dict]) -> list[int]:
    suggestions: list[int] = []
    actions = measure_actions(measure)
    by_id = _rule_prices(prices)
    for action in actions or [measure or ""]:
        normalized = normalize_text(action)
        # Exact source descriptions always win over broad keyword rules (e.g.
        # 60-minute glass, removal instead of installation, specific diameters).
        exact = [int(p['id']) for p in prices if normalize_text(p['description']) == normalized]
        if exact:
            if len(exact) == 1 and exact[0] not in suggestions:
                suggestions.append(exact[0])
            continue
        variant = None
        if re.search(r'draadglas|brandwerend glas|glas in (het kozijn|de deur)|zijlicht|bovenlicht|\bkoof\b|\bschacht\b', normalized) and not re.search(r'naden|aantonen', normalized):
            duration = re.search(r'\b(30|60)\s*min', normalized)
            if not duration:
                continue  # Unknown fire resistance is not a 30/60-minute default.
            minutes = int(duration[1])
            family = next((name for name in ('draadglas', 'koof', 'schacht') if name in normalized), 'glas')
            variant = {'draadglas': {30: 1, 60: 2}, 'glas': {30: 3, 60: 4},
                       'koof': {30: 35, 60: 34}, 'schacht': {30: 37, 60: 36}}[family][minutes]
        elif 'doorvoer' in normalized:
            matches = [(r'vwa|hyg', 38), (r'\brga\b', 39), (r'meterkast', 40), (r'<\s*22', 41), (r'>\s*22', 42), (r'verwarmingsbuis|verwarmingsbuiz', 43)]
            variants = [pid for pattern, pid in matches if re.search(pattern, normalized)]
            if len(variants) != 1:
                continue
            variant = variants[0]
        elif 'brandklep' in normalized:
            diameter = re.search(r'\b(125|250|315|400|500)\s*mm', normalized)
            if not diameter:
                continue
            variant = {125: 44, 250: 45, 315: 46, 400: 47, 500: 48}[int(diameter[1])]
        elif 'vluchtrouteaanduid' in normalized and 'verwijderen' in normalized:
            variant = 72
        elif 'rookmelder' in normalized and 'verplaatsen' in normalized:
            variant = 79
        elif 'rookmelder' in normalized and 'koppelbaar' in normalized:
            variant = 78
        elif 'vrijloopdranger' in normalized and 'aansturing' in normalized:
            control = 13 if 'bmi' in normalized else (14 if 'rf' in normalized else None)
            if control is None:
                continue
            for pid in [12, control]:
                if pid in by_id and by_id[pid]['id'] not in suggestions:
                    suggestions.append(by_id[pid]['id'])
            continue
        elif re.search(r'deur (zonder|met) kozijn vervangen', normalized):
            # Finish/glazing changes the source price substantially.
            if not re.search(r'schilderdeur|hpl|glasopening', normalized):
                continue
            variant = (30 if 'hpl' in normalized else 31 if 'glasopening' in normalized else 29) if 'met kozijn' in normalized else (27 if 'hpl' in normalized else 28 if 'glasopening' in normalized else 26)
        if variant is not None:
            if variant in by_id and by_id[variant]['id'] not in suggestions:
                suggestions.append(by_id[variant]['id'])
            continue
        for pattern, price_ids in PRICE_RULES:
            if re.search(pattern, normalized):
                for price_id in price_ids:
                    if price_id in by_id and by_id[price_id]['id'] not in suggestions:
                        suggestions.append(by_id[price_id]['id'])
                break
        # Explicit combined standard actions must not lose their second action.
        for condition, extra in [
            ('oplaadpunt' in normalized and 'verwijder' in normalized, 60),
            ('bewoners informeren' in normalized and 'scootmobiel' in normalized, 58),
            ('sluitkracht' in normalized and ('expanderende strips' in normalized or 'opschuimende strips' in normalized), 20),
        ]:
            if condition and extra in by_id and by_id[extra]['id'] not in suggestions:
                suggestions.append(by_id[extra]['id'])
        # No fuzzy fallback: similar wording can describe an opposite action.
    return suggestions


def measure_actions(measure):
    return [p.strip() for p in re.split(r'\n+|;', measure or '') if p.strip()]


def encode_price_links(ids, measure, prices, original_measure=None, original_ids=None):
    """Bind confirmed links to the final text and source identity, not just a row."""
    if original_measure is not None and normalize_text(measure) != normalize_text(original_measure) and list(ids) == list(original_ids or []):
        ids = suggest_price_ids(measure, prices)
    by_id = {int(p['id']): p for p in prices}
    unmatched = [a for a in measure_actions(measure) if not suggest_price_ids(a, prices)] if list(ids) == suggest_price_ids(measure, prices) else []
    return json.dumps({'version': 1, 'measure': normalize_text(measure), 'unmatched': unmatched, 'items': [
        {k: by_id[i][k] for k in ('id', 'description', 'unit')} for i in ids if i in by_id
    ]}, ensure_ascii=False)


def resolve_price_links(finding, prices):
    """Return source records and missing links; never drop an invalid selection."""
    value = finding.get('cost_items')
    try:
        raw = json.loads(value) if isinstance(value, str) and value else value
    except (ValueError, TypeError):
        return [], ['Opgeslagen prijskoppeling is onleesbaar; kies de prijsregel opnieuw.']
    measure = finding.get('maatregel') or ''
    by_id = {int(p['id']): p for p in prices}
    selected, missing = [], []
    if isinstance(raw, dict):
        if raw.get('measure') != normalize_text(measure):
            return [], [measure or 'Gewijzigde maatregel: controleer de prijskoppeling']
        for saved in raw.get('items', []):
            candidates = [p for p in prices if normalize_text(p['description']) == normalize_text(saved.get('description')) and p.get('unit', '') == saved.get('unit', '')]
            if len(candidates) == 1:
                selected.append(candidates[0])
            else:
                missing.append(saved.get('description') or measure)
        missing.extend(raw.get('unmatched', []))
        if not raw.get('items') and not missing:
            missing.append(measure or 'Geen prijsregel gekoppeld')
    elif parse_price_ids(raw):
        expected = suggest_price_ids(measure, prices)
        if measure.strip() and set(parse_price_ids(raw)) != set(expected):
            # Old records contain row numbers only: do not overwrite a possible
            # manual choice, but require reconfirmation instead of pricing blindly.
            return [], [f'Controleer bestaande prijskoppeling bij maatregel: {measure}']
        for pid in parse_price_ids(raw):
            if pid in by_id:
                selected.append(by_id[pid])
            else:
                missing.append(f'Ontbrekende prijsregel {pid}: {measure}')
    else:
        for action in measure_actions(measure) or ['Nog te bepalen']:
            ids = suggest_price_ids(action, prices)
            if ids:
                selected.extend(by_id[i] for i in ids)
            else:
                missing.append(action)
    return list({p['id']: p for p in selected}.values()), missing


def price_label(price: dict) -> str:
    unit = f" / {price['unit']}" if price.get("unit") else ""
    amount = f"{float(price['price']):,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return f"[{int(price['id']):03d}] {price['description']} — € {amount}{unit}"
