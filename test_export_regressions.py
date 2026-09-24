import io
import unittest
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

from report_export import build_report, DISCIPLINE_COLORS
from unit_prices import load_unit_prices, suggest_price_ids, encode_price_links, resolve_price_links, parse_price_ids
from cost_export import _cost_rows, build_cost_estimate

APP = Path(__file__).parent
PRICES = APP / 'data/eenheidsprijzen_2026.xlsx'
TEMPLATE = APP / 'templates/rapportage_brandveiligheid_template.docx'


class ExportRegressions(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.prices = load_unit_prices(PRICES)

    def test_all_source_descriptions_match_their_own_price(self):
        for price in self.prices:
            with self.subTest(id=price['id']):
                self.assertEqual(suggest_price_ids(price['description'], self.prices), [price['id']])

    def test_ambiguous_and_combined_measures(self):
        for measure in ['Draadglas vervangen met ... minuten WBDBO', 'Doorvoeren brandwerend afwerken',
                        'Ventilatiekanalen voorzien van brandklep', 'Vrijloopdranger toepassen, onderhouden en voorzien van aansturing.']:
            self.assertEqual(suggest_price_ids(measure, self.prices), [])
        self.assertEqual(suggest_price_ids('Draadglas moet vervangen worden met WBDBO van 60 minuten', self.prices), [2])
        self.assertEqual(suggest_price_ids('Vluchtrouteaanduidingen verwijderen conform voorschrift', self.prices), [72])
        self.assertEqual(set(suggest_price_ids('Oplaadpunten verwijderen en daarnaast bewoners informeren om geen scootmobielen te stallen in de vluchtwegen.', self.prices)), {58, 60})
        self.assertEqual(set(suggest_price_ids('Sluitkracht dranger aanpassen naar kracht 5 in combinatie met het toepassen van expanderende strips rondom de deur.', self.prices)), {11, 20})

    def test_final_measure_and_stable_identity(self):
        old = self.prices[0]['description']
        new = self.prices[1]['description']
        payload = encode_price_links([1], new, self.prices, old, [1])
        self.assertEqual(parse_price_ids(payload), [2])
        moved = [dict(p, id=p['id'] + 500) for p in self.prices]
        self.assertEqual(suggest_price_ids('Draadglas moet vervangen worden met WBDBO van 60 minuten', moved), [502])
        linked, missing = resolve_price_links({'maatregel': new, 'cost_items': payload}, moved)
        self.assertEqual([p['id'] for p in linked], [502])
        self.assertEqual(missing, [])
        linked, missing = resolve_price_links({'maatregel': old, 'cost_items': payload}, self.prices)
        self.assertFalse(linked)
        self.assertTrue(missing)
        empty = encode_price_links([], old, self.prices)
        self.assertFalse(resolve_price_links({'maatregel': old, 'cost_items': empty}, self.prices)[0])
        self.assertFalse(resolve_price_links({'maatregel': new, 'cost_items': '[1]'}, self.prices)[0])

    def test_invalid_and_partial_links_stay_visible(self):
        rows, count = _cost_rows([{'code_number': 1, 'cost_items': '[61,999]', 'aantal': '2'}], self.prices)
        self.assertEqual(len(rows), 2)
        self.assertEqual(count, 1)
        measure = self.prices[0]['description'] + '\n\nSpeciaal maatwerk zonder bronprijs'
        payload = encode_price_links([1], measure, self.prices)
        rows, count = _cost_rows([{'code_number': 2, 'cost_items': payload, 'maatregel': measure}], self.prices)
        self.assertEqual(count, 1)
        self.assertEqual(len(rows), 2)

    def test_word_colors_and_upright_footer(self):
        findings = [dict(finding_type='Maatregel', discipline=k, code_group='G', code_number=i+1, gebrek=k, maatregel='Controle') for i, k in enumerate(DISCIPLINE_COLORS)]
        output = build_report(TEMPLATE, {}, findings)
        doc = Document(io.BytesIO(output))
        tables = [t for t in doc.tables if t.rows[0].cells[0].text.startswith('Gebrek G.')]
        self.assertEqual(len(tables), 3)
        for table, fill in zip(tables, DISCIPLINE_COLORS.values()):
            for shade in table.rows[0]._tr.iter(qn('w:shd')):
                self.assertEqual(shade.get(qn('w:fill')), fill)
                self.assertIsNone(shade.get(qn('w:themeFill')))
        with ZipFile(io.BytesIO(output)) as z:
            xml = z.read('word/footer1.xml').decode()
            self.assertNotIn('rot="10800000"', xml)
            self.assertNotIn('flipH="1"', xml)
            self.assertNotIn('rotation:180', xml)
            self.assertIn('PAGE', xml)
            self.assertIn('TriaCon', xml)

    def test_excel_export_preserves_comments_and_flags_missing_price(self):
        output, unresolved = build_cost_estimate(APP / 'templates/kostenraming_2026_template.xlsx', PRICES, {}, [
            {'code_group': 'G', 'code_number': 1, 'maatregel': self.prices[1]['description'], 'aantal': '2'},
            {'code_group': 'W', 'code_number': 3, 'maatregel': 'Onbekend maatwerk', 'aantal': '1'},
        ])
        wb = load_workbook(io.BytesIO(output))
        s = wb['Blad1']
        self.assertEqual(s['C5'].value, self.prices[1]['description'])
        self.assertTrue(s['C6'].value.startswith('Nog koppelen:'))
        self.assertIsNone(s['D5'].value)
        self.assertIsNone(s['D6'].value)
        self.assertIn('NA()', s['G6'].value)
        self.assertEqual(unresolved, 1)
        wb.close()


if __name__ == '__main__':
    unittest.main()
