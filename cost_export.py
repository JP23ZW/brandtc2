from __future__ import annotations

import re
from copy import copy
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from unit_prices import load_unit_prices, resolve_price_links


DETAIL_ROWS = [5, 6, 7, 8, 9, 10, 12, 13, 14, 15]


def _copy_row_style(sheet: Worksheet, source_row: int, target_row: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _quantity(value) -> float:
    text = str(value or "").strip()
    # Accept both Dutch decimals (1,5 / 1.234,5) and a decimal point (1.5).
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return 1.0
    try:
        return float(match.group())
    except ValueError:
        return 1.0


def _cost_rows(findings: list[dict], prices: list[dict]) -> tuple[list[dict], int]:
    result = []
    unresolved = 0
    for finding in findings:
        if finding.get("finding_type", "Maatregel") != "Maatregel":
            continue
        valid_prices, missing = resolve_price_links(finding, prices)
        code = f"{finding.get('code_group', 'G')}.{int(finding.get('code_number') or 0):02d}"
        quantity = _quantity(finding.get("aantal"))
        for description in missing:
            unresolved += 1
            result.append(
                {
                    "code": code,
                    "description": f"Nog koppelen: {description}",
                    "quantity": quantity,
                    "unit": "",
                    "price": 0.0,
                }
            )
        for price in valid_prices:
            if float(price.get("price") or 0) == 0:
                unresolved += 1
            result.append(
                {
                    "code": code,
                    "description": price["description"],
                    "quantity": quantity,
                    "unit": price.get("unit", ""),
                    "price": float(price.get("price") or 0),
                }
            )
    return result, unresolved


def _fill_internal_price_sheet(workbook, prices: list[dict]) -> None:
    sheet = workbook["Blad2"]
    for row in range(1, max(sheet.max_row, len(prices)) + 1):
        for column in range(1, 4):
            sheet.cell(row, column).value = None
    for row, item in enumerate(prices, start=1):
        sheet.cell(row, 1).value = item["description"]
        sheet.cell(row, 2).value = float(item["price"])
        sheet.cell(row, 3).value = item["unit"]
    sheet.sheet_state = "hidden"
    if "EenheidsprijzenLijst" in workbook.defined_names:
        del workbook.defined_names["EenheidsprijzenLijst"]
    workbook.defined_names.add(
        DefinedName("EenheidsprijzenLijst", attr_text=f"'Blad2'!$A$1:$A${len(prices)}")
    )
    if "Blad3" in workbook.sheetnames:
        workbook["Blad3"].sheet_state = "hidden"


def build_cost_estimate(
    template_path: Path,
    price_path: Path,
    project: dict,
    findings: list[dict],
) -> tuple[bytes, int]:
    prices = load_unit_prices(price_path)
    rows, unresolved = _cost_rows(findings, prices)
    workbook = load_workbook(template_path, data_only=False, keep_links=False)
    sheet = workbook["Blad1"]
    _fill_internal_price_sheet(workbook, prices)

    extra_rows = max(0, len(rows) - len(DETAIL_ROWS))
    if extra_rows:
        sheet.insert_rows(16, extra_rows)
        for row in range(16, 16 + extra_rows):
            _copy_row_style(sheet, 15, row)

    detail_rows = [5, 6, 7, 8, 9, 10] + list(range(12, 16 + extra_rows))
    for row in detail_rows:
        for column in range(2, 9):
            sheet.cell(row, column).value = None

    sheet["B2"] = project.get("projectadres") or project.get("complexnaam") or "(straatnaam)"
    for row_number, item in zip(detail_rows, rows):
        sheet.cell(row_number, 2).value = item["code"]
        sheet.cell(row_number, 3).value = item["description"]
        # Opmerking TriaCon is intentionally reserved for manual completion in Excel.
        sheet.cell(row_number, 4).value = None
        sheet.cell(row_number, 5).value = float(item["quantity"])
        sheet.cell(row_number, 6).value = f'=IFERROR(VLOOKUP(C{row_number},\'Blad2\'!$A$1:$C${len(prices)},3,FALSE),"")'
        sheet.cell(row_number, 7).value = f'=IFERROR(VLOOKUP(C{row_number},\'Blad2\'!$A$1:$C${len(prices)},2,FALSE),NA())'
        sheet.cell(row_number, 8).value = f"=E{row_number}*G{row_number}"
        sheet.cell(row_number, 3).alignment = copy(sheet.cell(5, 3).alignment)
        sheet.cell(row_number, 3).alignment = sheet.cell(row_number, 3).alignment.copy(wrap_text=True)
        sheet.cell(row_number, 4).alignment = copy(sheet.cell(5, 4).alignment)
        sheet.cell(row_number, 4).alignment = sheet.cell(row_number, 4).alignment.copy(wrap_text=True)
        sheet.cell(row_number, 5).number_format = "0.00"
        sheet.cell(row_number, 7).number_format = '€ #,##0.00'
        sheet.cell(row_number, 8).number_format = '€ #,##0.00'

    price_validation = DataValidation(type="list", formula1="EenheidsprijzenLijst", allow_blank=True)
    price_validation.error = "Kies een werkzaamheid uit de eenheidsprijzenlijst 2026."
    price_validation.errorTitle = "Onbekende werkzaamheid"
    price_validation.prompt = "Selecteer een gekoppelde werkzaamheid."
    price_validation.promptTitle = "Eenheidsprijzen 2026"
    price_validation.showErrorMessage = True
    price_validation.showInputMessage = True
    sheet.add_data_validation(price_validation)
    for row in detail_rows:
        price_validation.add(sheet.cell(row, 3))

    first_group_total = 11
    second_group_total = 16 + extra_rows
    combined_total = 17 + extra_rows
    subtotal_row = 18 + extra_rows
    unforeseen_row = 19 + extra_rows
    abk_row = 20 + extra_rows
    tail_row = 21 + extra_rows
    vat_row = 22 + extra_rows
    total_row = 24 + extra_rows

    sheet.cell(first_group_total, 8).value = "=SUM(H5:H10)"
    sheet.cell(second_group_total, 8).value = f"=SUM(H12:H{15 + extra_rows})"
    sheet.cell(combined_total, 8).value = f"=H{first_group_total}+H{second_group_total}"
    sheet.cell(subtotal_row, 8).value = f"=H{combined_total}"
    sheet.cell(unforeseen_row, 7).value = f"=H{subtotal_row}*10%"
    sheet.cell(unforeseen_row, 8).value = f"=E{unforeseen_row}*G{unforeseen_row}"
    sheet.cell(abk_row, 7).value = f"=H{subtotal_row}*10%"
    sheet.cell(abk_row, 8).value = f"=G{abk_row}"
    sheet.cell(tail_row, 7).value = f"=H{subtotal_row}*10%"
    sheet.cell(tail_row, 8).value = f"=G{tail_row}"
    sheet.cell(vat_row, 7).value = f"=SUM(H{subtotal_row}:H{tail_row})*21%"
    sheet.cell(vat_row, 8).value = f"=G{vat_row}"
    sheet.cell(total_row, 8).value = f"=SUM(H{subtotal_row}:H{vat_row})"
    for row in [subtotal_row, unforeseen_row, abk_row, tail_row, vat_row, total_row]:
        sheet.cell(row, 7).number_format = '€ #,##0.00'
        sheet.cell(row, 8).number_format = '€ #,##0.00'

    sheet.print_area = f"$B$1:$H${total_row}"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), unresolved
