"""Import the supported TriaCon report layout without modifying the source."""
import io
import json
import re
import shutil
import uuid
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from docx_safety import repair_cover_anchors
from report_export import BUILDING_FIELDS, INSTALLATION_FIELDS, ORGANISATION_FIELDS, MATERIAL_FIELDS, DISCIPLINE_COLORS
import storage
from access_control import guard


def _text(element):
    return '\n'.join(''.join(n.text or '' for n in p.iter(qn('w:t'))) for p in element.iter(qn('w:p'))).strip()


def parse_report(content):
    repaired, repairs = repair_cover_anchors(content)
    doc = Document(io.BytesIO(repaired))
    if len(doc.tables) < 7 or [len(t.rows) for t in doc.tables[2:6]] != [13, 9, 8, 13]:
        raise ValueError('Dit rapport heeft niet de ondersteunde TriaCon-tabelstructuur. Er is niets geïmporteerd.')
    data, findings, drawings = {}, [], []
    assets = {}

    def picture(element):
        embeds = element.xpath('.//a:blip/@r:embed')
        if not embeds:
            return ''
        part = doc.part.related_parts[embeds[0]]
        key = str(part.partname)
        if key not in assets:
            # Canonical images, not executable attachments or arbitrary ZIP paths.
            with Image.open(io.BytesIO(part.blob)) as im:
                im.load()
                out = io.BytesIO()
                im.convert('RGB').save(out, 'PNG')
                assets[key] = out.getvalue()
        return key

    for table, keys in zip(doc.tables[2:6], [BUILDING_FIELDS, INSTALLATION_FIELDS, ORGANISATION_FIELDS, MATERIAL_FIELDS]):
        for row, key in zip(table.rows, keys):
            controls = row._tr.xpath('.//w:sdt')
            data[key] = _text(controls[-1]) if controls else row.cells[-1].text.strip()
    for p in doc.paragraphs:
        for label, key in [('Kenmerk:', 'kenmerk'), ('Datum:', 'report_date'), ('Versie:', 'version')]:
            if p.text.strip().startswith(label):
                data[key] = p.text.strip()[len(label):].strip()
    front = [r.cells[1].text.strip() for r in doc.tables[0].rows]
    if len(front) != 7:
        raise ValueError('De rapportgegevens zijn niet herkenbaar.')
    data.update(opdrachtgever=front[1].split('\n')[0], opdrachtgever_adres='\n'.join(front[1].split('\n')[1:]),
                opdrachtgever_contact=front[2], inspecteur=front[4], gecontroleerd=front[5],
                triacon_contact=front[6].split('\n')[0], triacon_email=front[6].split('E-mail:')[-1].strip() if 'E-mail:' in front[6] else '')
    data['import_source_context'] = {k: data.get(k, '') for k in storage.CONTEXT_KEYS if k in data}
    headings = {'Samenvatting': 'samenvatting', 'Algemene omschrijving complex': 'algemene_omschrijving',
                'Bezochte woningen': 'bezochte_woningen', 'Beperkingen': 'beperkingen', 'Conclusie': 'conclusie'}
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() in headings and p.style.name.lower().startswith(('heading', 'kop')):
            texts = []
            for following in doc.paragraphs[i + 1:]:
                if following.style.name.lower().startswith(('heading', 'kop')):
                    break
                if following.text.strip():
                    texts.append(following.text)
            data[headings[p.text.strip()]] = '\n'.join(texts).strip()
        full_text = ''.join(n.text or '' for n in p._p.iter(qn('w:t')))
        match = re.search(r'Voorafgaand aan de inspectie zijn (.*?) tekeningen ontvangen', full_text)
        if match:
            data['tekeningen_ontvangen'] = match[1]
        if p.text.strip() == 'Gelijkwaardigheidsoplossing':
            section = []
            for following in doc.paragraphs[i + 1:]:
                if following.style.name.lower().startswith(('heading', 'kop')):
                    break
                if following.text.strip():
                    section.append(following.text)
            data['gelijkwaardigheid'] = '\n'.join(section[1:])
        if p.text.startswith('Tekening ') and i + 1 < len(doc.paragraphs):
            asset = picture(doc.paragraphs[i + 1]._p)
            if asset:
                drawings.append({'name': p.text, 'image': asset})
    for cell, key in zip([c for r in doc.tables[6].rows for c in r.cells],
                         ['photo_voorgevel', 'photo_kopgevel', 'photo_achtergevel', 'photo_luchtfoto']):
        data[key] = picture(cell._tc)
    labels = {'tekeningnummer': 'tekeningnummer', 'bouwlaag': 'bouwlaag', 'ruimte(nummer)': 'ruimte',
              'ruimtenummer': 'ruimte', 'eis': 'eis', 'gebrek': 'gebrek', 'bevinding': 'gebrek',
              'aantal': 'aantal', 'afmeting': 'afmeting', 'maatregel': 'maatregel', 'conclusie': 'maatregel', 'opmerking': 'opmerking'}
    for table in doc.tables[7:]:
        title = table.rows[0].cells[0].text.strip()
        match = re.fullmatch(r'Gebrek\s+(.+)\.(\d+)', title)
        other = title == 'Foto tijdens inspectie'
        if not match and not other:
            continue
        finding = {'finding_type': 'Overige bevinding' if other else 'Maatregel',
                   'code_group': match[1] if match else 'O', 'code_number': int(match[2]) if match else len(findings) + 1,
                   'cost_items': '[]', 'discipline': 'Bouwkundig'}
        colors = table.rows[0]._tr.xpath('.//w:shd/@w:fill')
        finding['discipline'] = next((k for k, v in DISCIPLINE_COLORS.items() if v in colors), 'Bouwkundig')
        recognized = set()
        for row in table.rows[4:]:
            label = row.cells[0].text.strip().rstrip(':').lower()
            if label in labels:
                key = labels[label]
                finding[key] = row.cells[-1].text.strip()
                recognized.add(key)
        if 'gebrek' not in recognized:
            raise ValueError('Een gebrekentabel heeft een afwijkende indeling. Import afgebroken om gegevensverlies te voorkomen.')
        finding['photo_before'] = picture(table.rows[2].cells[0]._tc)
        finding['photo_after'] = picture(table.rows[2].cells[1]._tc)
        findings.append(finding)
    codes = [(f['code_group'], f['code_number']) for f in findings]
    if len(codes) != len(set(codes)):
        raise ValueError('Dubbele gebrekscodes in dit rapport; corrigeer deze vóór het inlezen.')
    warnings = ['Controleer de import vóór verdere rapportage. Onderwerp, richtlijn en prijskoppelingen staan niet afzonderlijk in dit Word-format en moeten zo nodig worden aangevuld.',
                'Vaste hoofdstukken en opmaak worden bij export opnieuw uit de app-template gebruikt; afwijkende vrije Word-opmaak en herstelregistratietabellen worden niet overgenomen. Het originele bestand blijft bewaard.']
    if drawings:
        warnings.append('Tekeningen worden als afbeelding met bestaande rode codes ingelezen. Bestaande markeringen zijn niet afzonderlijk verplaatsbaar; nieuwe markeringen kunnen wel worden toegevoegd.')
    if repairs:
        warnings.append('De bekende fout in het fotoanker is tijdens het lezen hersteld.')
    return {'data': data, 'findings': findings, 'drawings': drawings, 'assets': assets, 'warnings': warnings}


@guard('create_report', 'complex', 'complex_id')
def save_import(parsed, source, complex_id, title, user_id):
    """Stage private assets, then atomically commit all database records."""
    folder = storage.DATA_DIR / 'uploads' / ('import_' + uuid.uuid4().hex)
    folder.mkdir(parents=True, exist_ok=False)
    try:
        (folder / 'origineel.docx').write_bytes(source)
        paths = {}
        for i, (key, content) in enumerate(parsed['assets'].items()):
            path = folder / f'foto_{i}.png'
            path.write_bytes(content)
            paths[key] = str(path.relative_to(storage.DATA_DIR))
        data = dict(parsed['data'])
        data['import_original'] = str((folder / 'origineel.docx').relative_to(storage.DATA_DIR))
        for key in ('photo_voorgevel', 'photo_kopgevel', 'photo_achtergevel', 'photo_luchtfoto'):
            data[key] = paths.get(data.get(key), '')
        findings = [dict(f, photo_before=paths.get(f.get('photo_before'), ''), photo_after=paths.get(f.get('photo_after'), '')) for f in parsed['findings']]
        with storage.connect() as con:
            con.execute('BEGIN IMMEDIATE')
            stamp = storage.now_iso()
            clean = {k: v for k, v in data.items() if k not in storage.CONTEXT_KEYS}
            rid = con.execute('INSERT INTO reports(complex_id,title,status,data_json,created_by,updated_by,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)',
                              (complex_id, title.strip() or 'Ingelezen rapport', 'Concept', json.dumps(clean, ensure_ascii=False), user_id, user_id, stamp, stamp)).lastrowid
            cols = ['finding_type', 'code_group', 'code_number', 'discipline', 'onderwerp', 'tekeningnummer', 'bouwlaag', 'ruimte', 'eis', 'gebrek', 'aantal', 'afmeting', 'maatregel', 'opmerking', 'richtlijn', 'cost_items', 'photo_before', 'photo_after']
            for f in findings:
                con.execute(f"INSERT INTO findings(report_id,{','.join(cols)},created_by,updated_by,created_at,updated_at) VALUES({','.join('?' for _ in range(len(cols)+5))})",
                            [rid] + [f.get(k, '') for k in cols] + [user_id, user_id, stamp, stamp])
            for page, d in enumerate(parsed['drawings'], 1):
                with Image.open(io.BytesIO(parsed['assets'][d['image']])) as im:
                    width, height = im.size
                path = paths[d['image']]
                con.execute('INSERT INTO drawings(report_id,name,drawing_number,floor,page_number,original_path,image_path,width,height,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)',
                            (rid, d['name'], d['name'], '', page, path, path, width, height, stamp))
        return int(rid)
    except Exception:
        # Only the unique directory created by this operation is removed.
        shutil.rmtree(folder)
        raise
